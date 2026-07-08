import docx

doc_path = "/Users/nabhan/Downloads/SKRIPSI TA 2026/Tuntas Kilat/TA Nabhan 2026 - Rev26 (Final).docx"
doc = docx.Document(doc_path)
for i, p in enumerate(doc.paragraphs):
    if 'Tahap Observasi (Empathize)' in p.text:
        print(f"[{i-1}] {doc.paragraphs[i-1].text}")
        print(f"[{i}] {p.text}")
        print(f"[{i+1}] {doc.paragraphs[i+1].text}")
