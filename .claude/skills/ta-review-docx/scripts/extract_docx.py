#!/usr/bin/env python3
"""Ekstrak paragraf + tabel .docx ke teks yang bisa dibaca agent.

python-docx TIDAK memasukkan isi tabel ke `paragraphs`, jadi keduanya diekstrak
terpisah. Heading BAB sering paragraf Normal yang di-bold (bukan Word style),
maka pemetaan struktur pakai regex teks, bukan style name.

Pakai:  python3 extract_docx.py "naskah.docx" [out_prefix]
Output: <prefix>_paras.txt, <prefix>_tables.txt, dan ringkasan struktur ke stdout.
"""
import re
import sys

import docx  # pip install python-docx --break-system-packages -q


def main() -> None:
    if len(sys.argv) < 2:
        sys.exit("Pakai: python3 extract_docx.py <file.docx> [out_prefix]")
    src = sys.argv[1]
    prefix = sys.argv[2] if len(sys.argv) > 2 else "ta"
    d = docx.Document(src)

    with open(f"{prefix}_paras.txt", "w") as f:
        for i, p in enumerate(d.paragraphs):
            if p.text.strip():
                f.write(f"[{i}][{p.style.name}] {p.text}\n")

    with open(f"{prefix}_tables.txt", "w") as f:
        for ti, t in enumerate(d.tables):
            f.write(f"=== TABLE {ti} ({len(t.rows)}x{len(t.columns)}) ===\n")
            for r in t.rows:
                cells = [c.text.strip().replace("\n", " / ") for c in r.cells]
                f.write(" | ".join(cells) + "\n")

    bab = [
        (i, p.text.strip())
        for i, p in enumerate(d.paragraphs)
        if re.match(r"^(BAB\s+[IVX]+|DAFTAR PUSTAKA|LEMBAR|ABSTRAK)", p.text.strip())
    ]
    caps = [
        p.text.strip()
        for p in d.paragraphs
        if re.match(r"^Gambar \d+\.\d+", p.text.strip())
    ]
    print(f"paragraf berisi: {sum(1 for p in d.paragraphs if p.text.strip())}")
    print(f"tabel: {len(d.tables)} | caption gambar: {len(caps)}")
    print("struktur:")
    for i, t in bab:
        print(f"  [{i}] {t[:70]}")


if __name__ == "__main__":
    main()
